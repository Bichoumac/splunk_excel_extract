"""
xlsx_builder.py

Turns Splunk search results (output_mode=json_rows) into an .xlsx workbook
with openpyxl. This module has no dependency on Splunk so it can be unit
tested on its own.

    sheets = [SheetData(name, fields, rows), ...]
    data = build_workbook(sheets, ExportOptions())   # -> bytes
"""

import io
import re
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DEFAULT_SHEET = "Results"
MAX_SHEET_NAME = 31
MAX_CELL_CHARS = 32767                # Excel hard limit per cell
MAX_ROWS_PER_SHEET = 1048575          # 1,048,576 rows minus the header row
MAX_COL_WIDTH = 60
MIN_COL_WIDTH = 8
DATE_FORMAT = "yyyy-mm-dd hh:mm:ss"

# Integers or decimals with at most 15 significant integer digits and no
# leading zero: "42" -> 42, but "007" or "1234567890123456" stay text.
NUMBER_RE = re.compile(r"^-?(0|[1-9]\d{0,14})(\.\d+)?$")
SHEET_NAME_FORBIDDEN_RE = re.compile(r"[\\/?*\[\]:]")

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
WRAP = Alignment(wrap_text=True, vertical="top")


class ExportOptions(object):
    def __init__(self, keep_text=False, keep_internal=False, dates=True,
                 max_rows=MAX_ROWS_PER_SHEET):
        self.keep_text = keep_text
        self.keep_internal = keep_internal
        self.dates = dates
        self.max_rows = min(max_rows, MAX_ROWS_PER_SHEET)


class SheetData(object):
    def __init__(self, name, fields, rows):
        self.name = name
        self.fields = fields
        self.rows = rows


def make_sheet_names(requested, count):
    """Excel sheet names: max 31 chars, none of \\ / ? * [ ] : and unique
    (case-insensitive). Missing names become Results, Results 2, ..."""
    used = set()
    names = []
    for i in range(count):
        raw = requested[i] if i < len(requested) and requested[i] else (
            DEFAULT_SHEET if i == 0 else "%s %d" % (DEFAULT_SHEET, i + 1))
        name = SHEET_NAME_FORBIDDEN_RE.sub("_", str(raw))[:MAX_SHEET_NAME].strip("'") or DEFAULT_SHEET
        base = name
        suffix = 2
        while name.lower() in used:
            tail = " %d" % suffix
            suffix += 1
            name = base[:MAX_SHEET_NAME - len(tail)] + tail
        used.add(name.lower())
        names.append(name)
    return names


def parse_splunk_time(text):
    """Parse the ISO 8601 _time Splunk returns ("2026-10-08T14:05:00.000+02:00")
    and return a naive datetime holding the wall-clock time of that offset, i.e.
    the time zone configured for the Splunk user. Returns None if not parsable."""
    m = re.match(r"^(\d{4}-\d{2}-\d{2})[T ](\d{2}:\d{2}:\d{2})(\.\d+)?(Z|[+-]\d{2}:?\d{2})?$", text)
    if m:
        value = datetime.strptime(m.group(1) + " " + m.group(2), "%Y-%m-%d %H:%M:%S")
        if m.group(3):
            value = value.replace(microsecond=int(round(float(m.group(3)) * 1000000)) % 1000000)
        return value
    # Epoch seconds (e.g. when _time was rewritten as a number).
    if re.match(r"^\d{9,10}(\.\d+)?$", text):
        return datetime.fromtimestamp(float(text), timezone.utc).replace(tzinfo=None)
    return None


def to_value(value, field, opts):
    """Convert a raw Splunk value into a Python value for an Excel cell.
    Returns (value, is_date)."""
    if value is None:
        return None, False
    if isinstance(value, (list, tuple)):
        value = "\n".join("" if v is None else str(v) for v in value)   # multivalue field
    text = str(value)
    if text == "":
        return None, False

    if field == "_time" and opts.dates:
        parsed = parse_splunk_time(text)
        if parsed is not None:
            return parsed, True

    if not opts.keep_text and NUMBER_RE.match(text):
        return (float(text) if "." in text else int(text)), False

    text = ILLEGAL_CHARACTERS_RE.sub("", text)
    if len(text) > MAX_CELL_CHARS:
        text = text[:MAX_CELL_CHARS]
    return text, False


def kept_columns(fields, opts):
    return [i for i, name in enumerate(fields)
            if opts.keep_internal or not name.startswith("_") or name in ("_time", "_raw")]


def _text_cell(ws, value):
    cell = WriteOnlyCell(ws, value=value)
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"          # never write a formula, keep "=..." as text
    return cell


def write_sheet(wb, sheet, opts):
    ws = wb.create_sheet(title=sheet.name)
    cols = kept_columns(sheet.fields, opts)
    header = [sheet.fields[i] for i in cols]
    rows = sheet.rows[:opts.max_rows]

    # First pass: convert values and measure column widths.
    widths = [len(h) for h in header]
    converted = []
    for row in rows:
        out = []
        for c, idx in enumerate(cols):
            raw = row[idx] if idx < len(row) else None
            value, is_date = to_value(raw, sheet.fields[idx], opts)
            out.append((value, is_date))
            if is_date:
                length = len(DATE_FORMAT)
            elif value is None:
                length = 0
            else:
                length = len(str(value).split("\n", 1)[0])
            if length > widths[c]:
                widths[c] = length
        converted.append(out)

    # Write-only worksheets need dimensions, panes and filters set before rows.
    for c, w in enumerate(widths):
        ws.column_dimensions[get_column_letter(c + 1)].width = min(MAX_COL_WIDTH, max(MIN_COL_WIDTH, w + 2))
    if header:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(header)), len(converted) + 1)

    header_cells = []
    for h in header:
        cell = _text_cell(ws, h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        header_cells.append(cell)
    ws.append(header_cells)

    for out in converted:
        cells = []
        for value, is_date in out:
            cell = _text_cell(ws, value)
            if is_date:
                cell.number_format = DATE_FORMAT
            elif isinstance(value, str) and "\n" in value:
                cell.alignment = WRAP
            cells.append(cell)
        ws.append(cells)

    return len(converted)


def build_workbook(sheets, opts):
    """Build the workbook and return it as bytes."""
    wb = Workbook(write_only=True)
    for sheet in sheets:
        write_sheet(wb, sheet, opts)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

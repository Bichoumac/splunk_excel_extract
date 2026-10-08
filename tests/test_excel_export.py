"""
Unit tests for the excel_export endpoint, run outside Splunk:

    python3 -m unittest discover -s tests -v
"""

import base64
import io
import json
import os
import sys
import unittest
from datetime import datetime

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "src", "bin", "lib"))
sys.path.insert(0, os.path.join(ROOT, "src", "bin"))

from openpyxl import load_workbook  # noqa: E402

import excel_export_handler as handler  # noqa: E402
import xlsx_builder  # noqa: E402


class FakeSplunk(object):
    """Minimal splunkd: jobs with json_rows results, honouring offset/count and
    a maxresultrows cap."""

    def __init__(self, jobs=None, max_result_rows=3):
        self.jobs = jobs or {}
        self.max_result_rows = max_result_rows
        self.calls = []
        self.dispatched = []
        self.cancelled = []

    def request(self, method, path, getargs=None, postargs=None):
        self.calls.append((method, path, getargs, postargs))
        parts = path.strip("/").split("/")
        if method == "POST" and parts[-1] == "jobs":
            sid = "dispatched_%d" % len(self.dispatched)
            self.dispatched.append(postargs)
            self.jobs[sid] = {"fields": ["host", "count"], "rows": [["a", "1"], ["b", "2"]]}
            return 201, json.dumps({"sid": sid})
        if method == "POST" and parts[-1] == "control":
            self.cancelled.append(parts[-2])
            return 200, "{}"
        if parts[-1] == "results":
            job = self.jobs[parts[-2]]
            offset = int(getargs["offset"])
            count = min(int(getargs["count"]), self.max_result_rows)
            rows = job["rows"][offset:offset + count]
            return 200, json.dumps({"fields": job["fields"], "rows": rows})
        sid = parts[-1]
        if sid not in self.jobs:
            return 404, json.dumps({"messages": [{"type": "FATAL", "text": "Unknown sid."}]})
        job = self.jobs[sid]
        state = job.get("state", "DONE")
        return 200, json.dumps({"entry": [{"content": {
            "dispatchState": state, "isDone": state == "DONE", "isFailed": state == "FAILED",
            "messages": {"fatal": ["boom"]} if state == "FAILED" else {}}}]})


def call(fake, query, method="GET", settings=None):
    request = {"method": method, "query": [list(p) for p in query],
               "session": {"user": "alice", "authtoken": "tok"}}
    h = handler.ExcelExportHandler()
    return h.handle(json.dumps(request), client_factory=lambda token: fake,
                    settings=settings or handler.Settings())


def workbook(response):
    return load_workbook(io.BytesIO(base64.b64decode(response["payload_base64"])))


class BuilderTests(unittest.TestCase):
    def test_values_are_typed(self):
        opts = xlsx_builder.ExportOptions()
        self.assertEqual(xlsx_builder.to_value("42", "x", opts), (42, False))
        self.assertEqual(xlsx_builder.to_value("-3.5", "x", opts), (-3.5, False))
        self.assertEqual(xlsx_builder.to_value("007", "x", opts), ("007", False))
        self.assertEqual(xlsx_builder.to_value("1234567890123456", "x", opts), ("1234567890123456", False))
        self.assertEqual(xlsx_builder.to_value(["a", "b"], "x", opts), ("a\nb", False))
        self.assertEqual(xlsx_builder.to_value("", "x", opts), (None, False))
        self.assertEqual(xlsx_builder.to_value("bad\x01char", "x", opts), ("badchar", False))
        value, is_date = xlsx_builder.to_value("2026-10-08T14:05:03.250+02:00", "_time", opts)
        self.assertTrue(is_date)
        self.assertEqual(value, datetime(2026, 10, 8, 14, 5, 3, 250000))

    def test_keep_text_and_no_dates(self):
        opts = xlsx_builder.ExportOptions(keep_text=True, dates=False)
        self.assertEqual(xlsx_builder.to_value("42", "x", opts), ("42", False))
        self.assertEqual(xlsx_builder.to_value("2026-10-08T14:05:03.000+02:00", "_time", opts)[1], False)

    def test_sheet_names(self):
        names = xlsx_builder.make_sheet_names(["a/b", "A_B", "", "x" * 40], 4)
        self.assertEqual(names, ["a_b", "A_B 2", "Results 3", "x" * 31])

    def test_formula_is_written_as_text(self):
        data = xlsx_builder.build_workbook(
            [xlsx_builder.SheetData("S", ["v"], [["=SUM(A1:A9)"]])], xlsx_builder.ExportOptions())
        ws = load_workbook(io.BytesIO(data))["S"]
        self.assertEqual(ws["A2"].value, "=SUM(A1:A9)")
        self.assertEqual(ws["A2"].data_type, "s")


class HandlerTests(unittest.TestCase):
    def setUp(self):
        self.fake = FakeSplunk({
            "job1": {"fields": ["_time", "_cd", "host", "bytes", "user_id"], "rows": [
                ["2026-10-08T10:00:00.000+02:00", "1:2", "web-01", "100", "000123"],
                ["2026-10-08T10:01:00.000+02:00", "1:3", "web-02", "200", "000124"],
                ["2026-10-08T10:02:00.000+02:00", "1:4", "web-03", "300", "000125"],
                ["2026-10-08T10:03:00.000+02:00", "1:5", "web-04", "400", "000126"],
                ["2026-10-08T10:04:00.000+02:00", "1:6", ["a", "b"], "500", "000127"],
            ]},
            "job2": {"fields": [{"name": "host"}, {"name": "count"}], "rows": [["web-01", "5"]]},
            "empty": {"fields": [], "rows": []},
            "failed": {"fields": [], "rows": [], "state": "FAILED"},
        })

    def test_sid_export_reads_every_page(self):
        resp = call(self.fake, [("sid", "job1"), ("filename", "report"), ("timestamp", "false")])
        self.assertEqual(resp["status"], 200)
        self.assertIn('filename="report.xlsx"', resp["headers"]["Content-Disposition"])
        ws = workbook(resp)["Results"]
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(rows[0], ("_time", "host", "bytes", "user_id"))      # _cd dropped
        self.assertEqual(len(rows), 6)                                        # header + 5 (2 pages of 3)
        self.assertEqual(rows[1][0], datetime(2026, 10, 8, 10, 0))
        self.assertEqual(rows[1][2], 100)
        self.assertEqual(rows[1][3], "000123")
        self.assertEqual(rows[5][1], "a\nb")
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual(ws.auto_filter.ref, "A1:D6")

    def test_multi_sheet_and_postprocess(self):
        resp = call(self.fake, [("sid", "job1"), ("sid", "job2"), ("postprocess", ""),
                                ("postprocess", "| sort host"), ("sheet", "Events"), ("sheet", "Hosts"),
                                ("keep_internal", "true")])
        self.assertEqual(resp["status"], 200)
        wb = workbook(resp)
        self.assertEqual(wb.sheetnames, ["Events", "Hosts"])
        self.assertIn("_cd", [c.value for c in wb["Events"][1]])
        searches = [c[2].get("search") for c in self.fake.calls if c[1].endswith("/results")]
        self.assertIn("| sort host", searches)

    def test_search_mode_dispatches_and_cancels(self):
        resp = call(self.fake, [("search", "index=main | stats count by host"), ("earliest", "-24h"),
                                ("app", "my_app")])
        self.assertEqual(resp["status"], 200)
        self.assertEqual(self.fake.dispatched[0]["search"], "search index=main | stats count by host")
        self.assertEqual(self.fake.dispatched[0]["earliest_time"], "-24h")
        self.assertTrue(any(c[1] == "/servicesNS/alice/my_app/search/jobs" for c in self.fake.calls))
        self.assertEqual(self.fake.cancelled, ["dispatched_0"])

    def test_risky_commands_are_refused(self):
        resp = call(self.fake, [("search", "| makeresults | outputlookup x.csv")])
        self.assertEqual(resp["status"], 403)
        self.assertIn("outputlookup", json.loads(resp["payload"])["error"])
        resp = call(self.fake, [("sid", "job1"), ("postprocess", "| sendemail to=x@y")])
        self.assertEqual(resp["status"], 403)
        self.assertEqual(self.fake.dispatched, [])

    def test_search_can_be_disabled(self):
        resp = call(self.fake, [("search", "index=main")],
                    settings=handler.Settings({"allow_search": "false"}))
        self.assertEqual(resp["status"], 403)

    def test_errors(self):
        self.assertEqual(call(self.fake, [])["status"], 400)
        self.assertEqual(call(self.fake, [("sid", "../../etc")])["status"], 400)
        self.assertEqual(call(self.fake, [("sid", "job1"), ("search", "x")])["status"], 400)
        self.assertEqual(call(self.fake, [("sid", "unknown")])["status"], 404)
        self.assertEqual(call(self.fake, [("sid", "empty")])["status"], 404)
        self.assertEqual(call(self.fake, [("sid", "failed")])["status"], 400)
        self.assertEqual(call(self.fake, [("sid", "job1")], method="DELETE")["status"], 405)

    def test_post_form_and_unicode_filename(self):
        request = {"method": "POST", "query": [], "form": [["sid", "job2"], ["filename", "détail/erreurs"],
                                                             ["timestamp", "false"]],
                   "session": {"user": "alice", "authtoken": "tok"}}
        resp = handler.ExcelExportHandler().handle(json.dumps(request), client_factory=lambda t: self.fake,
                                                   settings=handler.Settings())
        self.assertEqual(resp["status"], 200)
        self.assertIn("filename*=UTF-8''d%C3%A9tail_erreurs.xlsx", resp["headers"]["Content-Disposition"])

    def test_settings_file(self):
        settings = handler.Settings.load(os.path.join(ROOT, "src"))
        self.assertTrue(settings.allow_search)
        self.assertIn("outputlookup", settings.risky_commands)
        self.assertEqual(settings.page_size, 50000)


if __name__ == "__main__":
    unittest.main()
